from fastapi import FastAPI
from pydantic import BaseModel
from typing import List, Dict
import re
import mmap

import fitz  
import pytesseract
from PIL import Image
from docx import Document
import openpyxl

app = FastAPI(title="pii-python-worker")

MAX_TEXT_BYTES = 512 * 1024
MAX_SAMPLES_PER_TYPE = 2


class InferRequest(BaseModel):
    path: str
    format: str  


class InferShmRequest(BaseModel):
    shm_path: str
    offset: int
    length: int
    format: str
    path: str


class Finding(BaseModel):
    category: str
    type: str
    count: int
    confidence: float
    masked_samples: List[str]


class InferResponse(BaseModel):
    findings: List[Finding]
    errors: List[str]

EMAIL_RE = re.compile(r"(?i)\b[a-z0-9._%+\-]+@[a-z0-9.\-]+\.[a-z]{2,}\b")
PHONE_RE = re.compile(r"(?:\+7|8)[\s\-\(]?\d{3}[\s\-\)]?\d{3}[\s\-]?\d{2}[\s\-]?\d{2}")
CARD_RE = re.compile(r"\b(?:\d[ -]*?){13,19}\b")
SNILS_RE = re.compile(r"\b\d{3}-?\d{3}-?\d{3}\s?\d{2}\b")
INN_RE = re.compile(r"\b\d{10}\b|\b\d{12}\b")

PASSPORT_RF_RE = re.compile(
    r"(?i)\bпаспорт(?:\s*гражданина\s*рф)?\b[^0-9]{0,40}(\d{2}\s?\d{2}\s?\d{6})\b"
)
BIK_RE = re.compile(r"\b04\d{7}\b")
BANK_ACCOUNT_RE = re.compile(r"\b\d{20}\b")
BIRTH_DATE_RE = re.compile(r"(?i)(?:дата\s*рождени[яе]|д\.р\.)[^0-9]{0,10}(\d{2}[.\-/]\d{2}[.\-/]\d{4})")
MRZ_RE = re.compile(r"\b[PIVAC][A-Z0-9<]{20,}\b")


def only_digits(s: str) -> str:
    return "".join(ch for ch in s if ch.isdigit())


def luhn_valid(num: str) -> bool:
    d = only_digits(num)
    if len(d) < 13 or len(d) > 19:
        return False
    s = 0
    alt = False
    for ch in reversed(d):
        n = int(ch)
        if alt:
            n *= 2
            if n > 9:
                n -= 9
        s += n
        alt = not alt
    return s % 10 == 0


def snils_valid(sn: str) -> bool:
    d = only_digits(sn)
    if len(d) != 11:
        return False
    num = d[:9]
    ctrl = int(d[9:])
    sm = sum(int(num[i]) * (9 - i) for i in range(9))
    if sm < 100:
        check = sm
    elif sm in (100, 101):
        check = 0
    else:
        check = sm % 101
        if check == 100:
            check = 0
    return check == ctrl


def inn_valid(inn: str) -> bool:
    d = only_digits(inn)
    if len(d) == 10:
        k = [2, 4, 10, 3, 5, 9, 4, 6, 8]
        c = sum(int(d[i]) * k[i] for i in range(9)) % 11 % 10
        return c == int(d[9])
    if len(d) == 12:
        k1 = [7, 2, 4, 10, 3, 5, 9, 4, 6, 8]
        c1 = sum(int(d[i]) * k1[i] for i in range(10)) % 11 % 10
        if c1 != int(d[10]):
            return False
        k2 = [3, 7, 2, 4, 10, 3, 5, 9, 4, 6, 8]
        c2 = sum(int(d[i]) * k2[i] for i in range(11)) % 11 % 10
        return c2 == int(d[11])
    return False


def mask_email(s: str) -> str:
    p = s.split("@")
    if len(p) != 2 or len(p[0]) < 2:
        return "***@***"
    return p[0][:2] + "***@" + p[1]


def mask_card(s: str) -> str:
    d = only_digits(s)
    if len(d) < 4:
        return "************"
    return "************" + d[-4:]


def clamp_text(text: str) -> str:
    if len(text.encode("utf-8", errors="ignore")) <= MAX_TEXT_BYTES:
        return text
    b = text.encode("utf-8", errors="ignore")[:MAX_TEXT_BYTES]
    return b.decode("utf-8", errors="ignore")


def detect_rules(text: str) -> List[Finding]:
    text = clamp_text(text)
    agg: Dict[str, Dict] = {}

    def add(tp: str, cat: str, sample: str, conf: float):
        if tp not in agg:
            agg[tp] = {"category": cat, "count": 0, "samples": [], "confidence": conf}
        agg[tp]["count"] += 1
        agg[tp]["confidence"] = max(agg[tp]["confidence"], conf)
        if len(agg[tp]["samples"]) < MAX_SAMPLES_PER_TYPE:
            agg[tp]["samples"].append(sample)

    for x in EMAIL_RE.findall(text):
        add("email", "ordinary", mask_email(x), 0.97)

    for _ in PHONE_RE.findall(text):
        add("phone", "ordinary", "+7******####", 0.95)

    for x in CARD_RE.findall(text):
        if luhn_valid(x):
            add("card_pan", "payment", mask_card(x), 0.98)

    for x in SNILS_RE.findall(text):
        if snils_valid(x):
            add("snils", "gov_id", "***-***-*** **", 0.99)

    for x in INN_RE.findall(text):
        if inn_valid(x):
            add("inn", "gov_id", "**********", 0.99)

    for _ in PASSPORT_RF_RE.findall(text):
        add("passport_rf", "gov_id", "**** ******", 0.92)

    for _ in BIK_RE.findall(text):
        add("bik", "payment", "04*******", 0.90)

    for _ in BANK_ACCOUNT_RE.findall(text):
        add("bank_account", "payment", "********************", 0.88)

    for _ in BIRTH_DATE_RE.findall(text):
        add("birth_date", "ordinary", "**.**.****", 0.86)

    for _ in MRZ_RE.findall(text):
        add("mrz", "gov_id", "P<********************", 0.93)

    out: List[Finding] = []
    for tp, v in agg.items():
        out.append(Finding(
            category=v["category"],
            type=tp,
            count=v["count"],
            confidence=v["confidence"],
            masked_samples=v["samples"],
        ))
    return out


def extract_pdf(path: str) -> str:
    text_parts = []
    doc = fitz.open(path)
    try:
        for page in doc:
            text_parts.append(page.get_text("text"))
    finally:
        doc.close()
    return "\n".join(text_parts)


def extract_docx(path: str) -> str:
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs)


def extract_xlsx(path: str) -> str:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts = []
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                vals = [str(v) for v in row if v is not None]
                if vals:
                    parts.append(" ".join(vals))
    finally:
        wb.close()
    return "\n".join(parts)


def extract_image_ocr(path: str) -> str:
    img = Image.open(path)
    try:
        return pytesseract.image_to_string(img, lang="rus+eng")
    finally:
        img.close()


@app.get("/health")
def health():
    return {"ok": True}


@app.post("/infer_shm", response_model=InferResponse)
def infer_shm(req: InferShmRequest):
    errors: List[str] = []
    text = ""

    try:
        with open(req.shm_path, "r+b") as f:
            mm = mmap.mmap(f.fileno(), 0, access=mmap.ACCESS_READ)
            try:
                chunk = mm[req.offset:req.offset + req.length]
            finally:
                mm.close()
        text = chunk.decode("utf-8", errors="ignore")
    except Exception as e:
        errors.append(f"shm read error: {e}")

    findings = detect_rules(text) if text else []
    return InferResponse(findings=findings, errors=errors)


@app.post("/infer", response_model=InferResponse)
def infer(req: InferRequest):
    errors: List[str] = []
    text = ""
    fmt = req.format.lower().strip()
    path = req.path

    try:
        if fmt == "pdf":
            text = extract_pdf(path)
        elif fmt == "docx":
            text = extract_docx(path)
        elif fmt in ("xlsx", "xlsm"):
            text = extract_xlsx(path)
        elif fmt in ("tif", "tiff", "jpg", "jpeg", "png", "gif"):
            text = extract_image_ocr(path)
        elif fmt == "mp4":
            errors.append("unsupported_media: mp4")
        else:
            errors.append(f"unsupported for python worker: {fmt}")
    except Exception as e:
        errors.append(str(e))

    findings = detect_rules(text) if text else []
    return InferResponse(findings=findings, errors=errors)